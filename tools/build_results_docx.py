"""Results sections as a .docx, in the layout the manuscript already uses.

This reproduces the structure of the existing "4.1 TCGA Validation Results /
4.2 PAIP Validation Results / 4.3 SurGen External Test Results" sections -
per-model tables for TCGA, per-aggregation tables for the validation cohorts,
and the "Summary of Best Results" table - but fills every cell from the
provenance-stamped ``result_*.json`` files rather than by hand.

Sections emitted
----------------
    4.1  TCGA Validation Results            TCGA-CV     (internal 4-fold CV)
    4.2  PAIP Validation Results            PAIP-EV     (TCGA-trained -> PAIP)
    4.3  SurGen External Test Results       SurGen-EV   (TCGA-trained -> SurGen)
    4.4  SurGen Internal Cross-Validation   SurGen-CV   (SurGen-trained, 4-fold)

4.4 is new: the manuscript predates the SurGen CV sweep, and 4.3 previously
carried H-Optimus-1 only because that was all that had been extracted. Both now
cover all five foundation models.

Selection rule
--------------
"Best Aggregation" and "Best Adaptation" are chosen by ``--select-by``, default
AUROC - the project treats AUROC as primary everywhere because it is
threshold-free (see runners/thresholds.py). Pass ``--select-by BalAcc`` to match
the selection used in the older hand-built summary table.

Run:  python tools/build_results_docx.py [--out FILE.docx] [--select-by AUROC]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
SLIDE_CLS = REPO_ROOT / "slide_classification"
OUT_DOCX = REPO_ROOT / "Results_Sections_UPDATED.docx"
sys.path.insert(0, str(REPO_ROOT / "tools"))

from build_report import discover, _rows_from, SLIDE_LEVEL_ENCODERS  # noqa: E402

FOUNDATION_MODELS = ["Conch1_5", "UNI2", "H-Optimus-1", "ConchV1", "Virchow2"]

MODEL_HEADING = {"Conch1_5": "Conch1_5", "UNI2": "UNI-2h",
                 "H-Optimus-1": "H-Optimus-1", "ConchV1": "Conch_V1",
                 "Virchow2": "Virchow2"}

MODEL_SHORT = {"Conch1_5": "Conch 1.5", "UNI2": "UNI-2h",
               "H-Optimus-1": "H-Optimus-1", "ConchV1": "Conch V1",
               "Virchow2": "Virchow2"}

AGG_METHODS = ["Averaging", "Caption_based_aggregation",
               "Caption_based_aggregation_15_classes", "Tissue_Type_Clustering"]

AGG_TITLE = {
    "Averaging": "Averaging Aggregation (1x feature dim)",
    "Caption_based_aggregation": "Caption-Based Aggregation (14x feature dim)",
    "Caption_based_aggregation_15_classes":
        "Caption-Based Aggregation, 15 classes (15x feature dim)",
    "Tissue_Type_Clustering": "Tissue-Type Clustering Aggregation (9x feature dim)",
}

AGG_SHORT = {"Averaging": "Averaging",
             "Caption_based_aggregation": "Caption-based clustering",
             "Caption_based_aggregation_15_classes": "Caption-based clustering (15)",
             "Tissue_Type_Clustering": "Tissue-type clustering"}

AGG_SUMMARY = {"Averaging": "Averaging",
               "Caption_based_aggregation": "Caption",
               "Caption_based_aggregation_15_classes": "Caption-15",
               "Tissue_Type_Clustering": "Tissue-type"}

CLASSIFIERS = ["lin", "ann", "knn", "proto", "rf"]
CLF_LONG = {"lin": "Linear", "ann": "ANN", "knn": "KNN",
            "proto": "ProtoNet", "rf": "r_forest"}
CLF_SHORT = {"lin": "LR", "ann": "ANN", "knn": "KNN",
             "proto": "ProtoNet", "rf": "RF"}

# thresholds.py rule 2: internal experiments (TCGA-CV, PAIP-IV, SurGen-CV)
# report at 0.5, which is the flat "default" row. PAIP-IV also carries an
# `at_tau_train` operating point; that is supplementary, not the headline.
PRIMARY_VARIANT = {"TCGA-CV": "default", "SurGen-CV": "default",
                   "PAIP-IV": "default", "PAIP-EV": "tcga_full",
                   "SurGen-EV": "tcga_full"}

INTRO = {
    "PAIP-EV": ("After training on TCGA, models were validated on the PAIP cohort "
                "across all aggregation strategies. No PAIP data enters training and "
                "the decision threshold is fixed on TCGA (Youden's J on pooled "
                "out-of-fold TCGA-CV probabilities), so the operating point is never "
                "tuned on the target cohort. Confusion matrices are [[TN, FP], [FN, TP]]."),
    "SurGen-EV": ("The SurGen cohort (two surgical sets, SR1482 and SR386, CZI format "
                  "at 40x; EBI accession S-BIAD1285) serves as a fully external test. "
                  "All five foundation models are now available. SurGen is markedly "
                  "imbalanced - 60 MSI-H among 622 slides (9.6%) - which strongly "
                  "shapes threshold-dependent metrics. Models are TCGA-trained and "
                  "scored at the TCGA-fixed threshold. Confusion matrices are "
                  "[[TN, FP], [FN, TP]]."),
    "SurGen-CV": ("Internal 4-fold cross-validation within SurGen, with folds built at "
                  "case level so no case is split across folds. Each fold serves as "
                  "test once, the next fold as validation, and the remaining two as "
                  "training. Reported at a fixed 0.5 threshold. Because prevalence is "
                  "9.6%, the majority-class accuracy baseline is 0.9035 - accuracy "
                  "should not be read as evidence of performance here."),
    "PAIP-IV": ("Internal validation on the PAIP provider-defined split (47 train / "
                "31 test), single split, no averaging."),
}


def pct(v) -> str:
    return "-" if pd.isna(v) else f"{float(v) * 100:.2f}%"


def dec(v) -> str:
    return "-" if pd.isna(v) else f"{float(v):.4f}"


def cm(v) -> str:
    return "-" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v)


def load_table() -> pd.DataFrame:
    rows: List[dict] = []
    for rec in discover(SLIDE_CLS):
        rows.extend(_rows_from(rec))
    if not rows:
        raise SystemExit("no result files found - run the experiments first")
    return pd.DataFrame(rows)


def experiment_slice(df: pd.DataFrame, exp: str) -> pd.DataFrame:
    sub = df[df["Experiment"] == exp]
    variant = PRIMARY_VARIANT.get(exp)
    if variant and variant in set(sub["Variant"]):
        sub = sub[sub["Variant"] == variant]
    return sub


# ------------------------------------------------------------------ docx helpers

def heading(doc, text, size, bold=True, space_before=14):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


def body(doc, text, size=10.5, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    return p


def table(doc, headers: List[str], rows: List[List[str]],
          bold_row: Optional[int] = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
            if bold_row is not None and ri == bold_row:
                run.bold = True
    doc.add_paragraph()
    return t


def best_row(sub: pd.DataFrame, metric: str) -> Optional[pd.Series]:
    s = sub.dropna(subset=[metric])
    return None if s.empty else s.loc[s[metric].idxmax()]


# ------------------------------------------------------------------ sections

def section_tcga(doc, df: pd.DataFrame, number: str, select_by: str):
    exp = "TCGA-CV"
    sub = experiment_slice(df, exp)
    if sub.empty:
        return
    heading(doc, f"{number} TCGA Validation Results", 15, space_before=0)
    n = sorted(sub["N_test"].dropna().unique().astype(int))
    body(doc, f"Internal 4-fold cross-validation on TCGA (N = "
              f"{'/'.join(str(v) for v in n)}). For each foundation model the table "
              f"reports its best aggregation, selected by {select_by}. Confusion "
              f"matrices are [[TN, FP], [FN, TP]].")

    for idx, model in enumerate(FOUNDATION_MODELS, start=1):
        ms = sub[(sub["Model"] == model) & (sub["Method"].isin(AGG_METHODS))]
        if ms.empty:
            continue
        scores = {a: g[select_by].max() for a, g in ms.groupby("Method")
                  if g[select_by].notna().any()}
        if not scores:
            continue
        best_agg = max(scores, key=scores.get)
        heading(doc, f"{idx}. {MODEL_HEADING.get(model, model)}", 12)
        body(doc, f"Best Aggregation: {AGG_SHORT.get(best_agg, best_agg)}")

        block = ms[ms["Method"] == best_agg]
        rows, best_i, best_v = [], None, -1.0
        for i, clf in enumerate(CLASSIFIERS):
            r = block[block["Classifier"] == clf]
            if r.empty:
                continue
            r = r.iloc[0]
            if pd.notna(r[select_by]) and float(r[select_by]) > best_v:
                best_v, best_i = float(r[select_by]), len(rows)
            rows.append([CLF_LONG[clf], pct(r.Acc), pct(r.BalAcc),
                         pct(r.MacroF1), pct(r.AUROC), cm(r.ConfMatrix)])
        table(doc, ["Classifier", "Acc", "Bal Acc", "Macro-F1", "AUROC", "CM"],
              rows, bold_row=best_i)

    # ---- Summary of Best Results
    heading(doc, "Summary of Best Results", 12)
    body(doc, f"One row per foundation model and aggregation method; the adaptation "
              f"shown is the classifier with the highest {select_by} for that pair.",
         italic=True)
    rows = []
    for method in AGG_METHODS:
        for model in FOUNDATION_MODELS:
            block = sub[(sub["Method"] == method) & (sub["Model"] == model)]
            r = best_row(block, select_by)
            if r is None:
                continue
            rows.append([MODEL_SHORT.get(model, model),
                         AGG_SUMMARY.get(method, method),
                         CLF_SHORT[r.Classifier],
                         f"{r.BalAcc * 100:.2f}", f"{r.AUROC * 100:.2f}"])
    for enc in sorted(SLIDE_LEVEL_ENCODERS):
        block = sub[sub["Method"] == enc]
        r = best_row(block, select_by)
        if r is None:
            continue
        label = "Conch with TITAN" if enc == "TITAN" else enc
        rows.append([label, enc, CLF_SHORT[r.Classifier],
                     f"{r.BalAcc * 100:.2f}", f"{r.AUROC * 100:.2f}"])
    table(doc, ["Models", "Aggregation", "Best Adaptation",
                "Balanced Acc", "AUROC"], rows)


def has_corrected(sub: pd.DataFrame) -> bool:
    """True only when the run carried --threshold-mode=corrected.

    Absent on a frozen run, so the tables keep exactly their published shape.
    """
    return ("BalAcc_corrected" in sub.columns
            and sub["BalAcc_corrected"].notna().any())


def has_promoted(sub: pd.DataFrame) -> bool:
    """True when the run carried --threshold-mode=promoted.

    Gated on ``Threshold_scheme`` rather than ``BalAcc_corrected``: under
    ``promoted`` the corrected values are moved INTO the headline metric fields
    for the promoted heads only, so ``BalAcc_corrected`` is never written and
    gating on it would silently render a mixed-scheme table as if it were
    single-scheme.
    """
    return ("Threshold_scheme" in sub.columns
            and sub["Threshold_scheme"].notna().any())


def scheme_cell(row) -> str:
    """Which threshold rule produced this row's headline numbers.

    Non-promoted heads carry no scheme, and are still at tau_TCGA - naming that
    explicitly is the point of the column: two schemes coexist in one table and
    the reader must be able to tell them apart.
    """
    s = row.get("Threshold_scheme")
    return "frozen (tau_TCGA)" if (s is None or pd.isna(s)) else str(s)


def status_cell(row, key: str = "Status_corrected") -> str:
    """Report a dead/saturated operating point instead of a bare 0.5000."""
    st = row.get(key)
    return "-" if (st is None or pd.isna(st)) else str(st)


def has_ci(sub: pd.DataFrame) -> bool:
    """Only iv_runner bootstraps, so CI columns exist for PAIP-IV alone.

    Checked from the data rather than hardcoded by experiment name, so the
    columns appear automatically if the other runners gain bootstrapping.
    """
    return all(c in sub.columns and sub[c].notna().any()
               for c in ("BalAcc_CI", "AUROC_CI"))


def ci(v) -> str:
    return "-" if (v is None or (isinstance(v, float) and pd.isna(v))) else str(v)


def section_by_aggregation(doc, df: pd.DataFrame, number: str, title: str,
                           exp: str, select_by: str):
    sub = experiment_slice(df, exp)
    if sub.empty:
        return
    heading(doc, f"{number} {title}", 15, space_before=0)
    if exp in INTRO:
        body(doc, INTRO[exp])
    show_ci = has_ci(sub)
    show_corr = has_corrected(sub)
    show_prom = has_promoted(sub)
    if show_ci:
        body(doc, "Bracketed values are 95% confidence intervals from 1000 "
                  "bootstrap resamples of the test split.", italic=True)
    if show_prom:
        body(doc, "This table mixes two threshold schemes, so every row names "
                  "its own. Scheme = refit_tau_k35 (KNN, re-scored at k=35) or "
                  "quantile_rate_matched (RF); rows reading frozen (tau_TCGA) "
                  "are unchanged from the published zero-shot result. Bal-Acc "
                  "and AUROC are the headline values under that row's scheme; "
                  "@prev columns give the previous frozen values so the change "
                  "is visible. Status flags an operating point that is dead (no "
                  "slide predicted positive) or saturated (all of them) - there "
                  "a Bal-Acc near 0.50 reflects the threshold, not the ranking.",
             italic=True)
    if show_corr:
        body(doc, "Bal-Acc @frozen is the strict zero-shot operating point "
                  "(tau_TCGA, unchanged). Bal-Acc @corrected uses the refit "
                  "threshold at k=35 for KNN and the rate-matched quantile for "
                  "LR/ANN/ProtoNet/RF. Status flags an operating point that is "
                  "dead (no slide predicted positive) or saturated (all of "
                  "them) - there a Bal-Acc near 0.50 reflects the threshold, "
                  "not the ranking.", italic=True)

    part = 1
    for method in AGG_METHODS:
        block = sub[sub["Method"] == method]
        if block.empty:
            continue
        n = sorted(block["N_test"].dropna().unique().astype(int))
        heading(doc, f"{number}.{part}  {AGG_TITLE.get(method, method)}"
                     f"   (N = {'/'.join(str(v) for v in n)})", 11.5)
        part += 1
        rows, best_i, best_v = [], None, -1.0
        for model in FOUNDATION_MODELS:
            for clf in CLASSIFIERS:
                r = block[(block["Model"] == model) & (block["Classifier"] == clf)]
                if r.empty:
                    continue
                r = r.iloc[0]
                if pd.notna(r[select_by]) and float(r[select_by]) > best_v:
                    best_v, best_i = float(r[select_by]), len(rows)
                row = [model, CLF_SHORT[clf]]
                if show_prom:
                    row.append(scheme_cell(r))
                row += [dec(r.Acc), dec(r.BalAcc)]
                if show_ci:
                    row.append(ci(r.get("BalAcc_CI")))
                if show_corr:
                    row.append(dec(r.get("BalAcc_corrected")))
                if show_prom:
                    row.append(dec(r.get("BalAcc_prev_frozen")))
                row += [dec(r.MacroF1), dec(r.AUROC)]
                if show_ci:
                    row.append(ci(r.get("AUROC_CI")))
                if show_corr:
                    row.append(dec(r.get("AUROC_corrected")))
                    row.append(status_cell(r))
                if show_prom:
                    row.append(dec(r.get("AUROC_prev_frozen")))
                    row.append(status_cell(r, "Status"))
                    row.append(status_cell(r, "Status_prev_frozen"))
                row.append(cm(r.ConfMatrix))
                rows.append(row)
        cols = ["Model", "Clf"]
        if show_prom:
            cols.append("Scheme")
        cols += ["Acc", "Bal-Acc"]
        if show_ci:
            cols.append("Bal-Acc 95% CI")
        if show_corr:
            cols.append("Bal-Acc @corr")
        if show_prom:
            cols.append("Bal-Acc @prev")
        cols += ["Mac-F1", "AUROC"]
        if show_ci:
            cols.append("AUROC 95% CI")
        if show_corr:
            cols += ["AUROC @corr", "Status"]
        if show_prom:
            cols += ["AUROC @prev", "Status", "Status @prev"]
        cols.append("Avg CM")
        table(doc, cols, rows, bold_row=best_i)

    for enc in ("TITAN", "PRISM"):
        block = sub[sub["Method"] == enc]
        if block.empty:
            continue
        models = sorted(block["Model"].unique())
        n = sorted(block["N_test"].dropna().unique().astype(int))
        note = (f"{enc} Aggregation ({', '.join(models)} only)"
                if enc == "TITAN" else f"{enc} Aggregation (slide-level embedding)")
        heading(doc, f"{number}.{part}  {note}   (N = {'/'.join(str(v) for v in n)})",
                11.5)
        part += 1
        rows, best_i, best_v = [], None, -1.0
        for clf in CLASSIFIERS:
            r = block[block["Classifier"] == clf]
            if r.empty:
                continue
            r = r.iloc[0]
            if pd.notna(r[select_by]) and float(r[select_by]) > best_v:
                best_v, best_i = float(r[select_by]), len(rows)
            row = [CLF_SHORT[clf]]
            if show_prom:
                row.append(scheme_cell(r))
            row += [dec(r.Acc), dec(r.BalAcc)]
            if show_ci:
                row.append(ci(r.get("BalAcc_CI")))
            if show_corr:
                row.append(dec(r.get("BalAcc_corrected")))
            if show_prom:
                row.append(dec(r.get("BalAcc_prev_frozen")))
            row += [dec(r.MacroF1), dec(r.AUROC)]
            if show_ci:
                row.append(ci(r.get("AUROC_CI")))
            if show_corr:
                row.append(dec(r.get("AUROC_corrected")))
                row.append(status_cell(r))
            if show_prom:
                row.append(dec(r.get("AUROC_prev_frozen")))
                row.append(status_cell(r, "Status"))
                row.append(status_cell(r, "Status_prev_frozen"))
            row.append(cm(r.ConfMatrix))
            rows.append(row)
        cols = ["Classifier"]
        if show_prom:
            cols.append("Scheme")
        cols += ["Acc", "Bal-Acc"]
        if show_ci:
            cols.append("Bal-Acc 95% CI")
        if show_corr:
            cols.append("Bal-Acc @corr")
        if show_prom:
            cols.append("Bal-Acc @prev")
        cols += ["Mac-F1", "AUROC"]
        if show_ci:
            cols.append("AUROC 95% CI")
        if show_corr:
            cols += ["AUROC @corr", "Status"]
        if show_prom:
            cols += ["AUROC @prev", "Status", "Status @prev"]
        cols.append("Avg CM")
        table(doc, cols, rows, bold_row=best_i)


def build(out: Path, select_by: str, sections: List[str]) -> Path:
    df = load_table()
    doc = Document()
    heading(doc, "Results - regenerated from the current result files", 17,
            space_before=0)
    stamp = df["Timestamp"].max() if "Timestamp" in df.columns else "?"
    commits = sorted(set(df["GitCommit"].dropna())) if "GitCommit" in df else []
    body(doc, f"Every value is read from the provenance-stamped result_*.json files. "
              f"Latest run: {stamp}. Git commit(s): {', '.join(commits) or '?'}. "
              f"Best row in each table (by {select_by}) is bold.", italic=True)

    plan = [("4.1", "TCGA Validation Results", "TCGA-CV"),
            ("4.2", "PAIP Validation Results", "PAIP-EV"),
            ("4.3", "SurGen External Test Results", "SurGen-EV"),
            ("4.4", "SurGen Internal Cross-Validation Results", "SurGen-CV"),
            ("4.5", "PAIP Internal Validation Results", "PAIP-IV")]

    for num, title, exp in plan:
        if exp not in sections:
            continue
        if exp == "TCGA-CV":
            section_tcga(doc, df, num, select_by)
        else:
            section_by_aggregation(doc, df, num, title, exp, select_by)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, default=OUT_DOCX)
    ap.add_argument("--select-by", default="AUROC", choices=["AUROC", "BalAcc"])
    ap.add_argument("--sections", default="TCGA-CV,PAIP-EV,SurGen-EV,SurGen-CV",
                    help="comma list; add PAIP-IV to include it")
    args = ap.parse_args()
    path = build(args.out, args.select_by, args.sections.split(","))
    size = path.stat().st_size / 1024
    print(f"wrote {path}  ({size:.0f} KB, selection metric: {args.select_by})")


if __name__ == "__main__":
    main()
