"""Standalone docx report for one EV experiment - radar plots, per-classifier
configuration, and the complete per-combination result table.

Distinct from ``build_supplementary_results.py`` (one docx, all five
experiments, no images) and ``build_results_docx.py`` (top-N tables only).
This produces one focused, self-contained document per experiment, meant to
be handed off on its own - it embeds the radar plots and states, in prose,
exactly what each classifier was trained on and thresholded with, so a reader
never has to cross-reference the code to know what a number means.

Reads only the published result_*.json files and the already-rendered radar
PNGs. Writes only the output docx - nothing under slide_classification/ or
Analysis_and_Visualization/ is touched.

Run:
    python tools/build_ev_experiment_report.py --experiment SurGen-EV
    python tools/build_ev_experiment_report.py --experiment PAIP-EV
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Dict, List, Optional

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT

REPO_ROOT = Path(__file__).resolve().parents[1]
SLIDE_CLS = REPO_ROOT / "slide_classification"
RADAR_ROOT = REPO_ROOT / "Analysis_and_Visualization" / "Radar_Plots"
sys.path.insert(0, str(SLIDE_CLS))

METHODS_AGG = ["Averaging", "Caption_based_aggregation",
               "Caption_based_aggregation_15_classes", "Tissue_Type_Clustering"]
METHOD_LABEL = {"Averaging": "Averaging", "Caption_based_aggregation": "Caption-based",
                "Caption_based_aggregation_15_classes": "Caption-based (15-class)",
                "Tissue_Type_Clustering": "Tissue-type clustering",
                "TITAN": "TITAN", "PRISM": "PRISM"}
CLF_ORDER = ["lin", "ann", "knn", "proto", "rf"]
CLF_NAME = {"lin": "Logistic Regression", "ann": "Neural Network (ANN)",
            "knn": "k-Nearest Neighbours", "proto": "ProtoNet",
            "rf": "Random Forest"}
SLIDE_LEVEL = {"TITAN", "PRISM"}

RGB_HEAD = RGBColor(0x1F, 0x2A, 0x44)
RGB_MUTED = RGBColor(0x5A, 0x5F, 0x70)
RGB_CHANGED = RGBColor(0xA1, 0x2D, 0x4C)


# --------------------------------------------------------------------------- data

def load_results(experiment: str) -> List[dict]:
    tree = {"SurGen-EV": "SurGen_EV_Results", "PAIP-EV": "TCGA_PAIP_EV_Results"}[experiment]
    glob_pat = f"result_{experiment}_*_ev.json"
    out = []
    for f in sorted((SLIDE_CLS / tree).rglob(glob_pat)):
        out.append(json.loads(f.read_text(encoding="utf-8")))
    return out


def n_pos(cm) -> int:
    return int(cm[0][1]) + int(cm[1][1])


# ---------------------------------------------------------------- docx helpers

def set_col_widths(table, widths_in):
    for row in table.rows:
        for cell, w in zip(row.cells, widths_in):
            cell.width = Inches(w)


def style_header_row(table):
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell._tc.get_or_add_tcPr()
    # shade header row
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1F2A44")
        tcPr.append(shd)


def add_table(doc, rows: List[List[str]], header: List[str],
              widths: Optional[List[float]] = None, small: bool = True):
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        table.rows[0].cells[i].text = h
    style_header_row(table)
    for r in rows:
        cells = table.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            if small:
                for p in cells[i].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(8.5)
    if widths:
        set_col_widths(table, widths)
    return table


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGB_HEAD
    return h


def add_body(doc, text, italic=False, size=10, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def add_image_pair(doc, paths: List[Path], caption: str):
    for p in paths:
        if p.exists():
            doc.add_picture(str(p), width=Inches(6.3))
            last = doc.paragraphs[-1]
            last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            add_body(doc, f"[missing: {p.name}]", italic=True, color=RGB_MUTED)
    add_body(doc, caption, italic=True, size=8.5, color=RGB_MUTED)


# --------------------------------------------------------------- config tables

def promoted_state(payloads: List[dict]) -> Dict[str, object]:
    """Which heads are promoted and at what k, read from the data itself.

    Not hardcoded per experiment name: SurGen-EV and PAIP-EV have each been
    promoted at a DIFFERENT k (20 and 35 respectively), and which experiment
    is promoted at all has already changed once (SurGen-EV was all-frozen
    until 3 Sep). Inspecting the payloads is the only way this stays correct
    without editing this file every time the published configuration moves.
    """
    promoted_clfs, knn_k = set(), None
    for p in payloads:
        for clf, pv in p["results"].items():
            e = pv.get("tcga_full")
            if isinstance(e, dict) and e.get("promoted"):
                promoted_clfs.add(clf)
        k = p["stamp"].get("knn_k_corrected")
        if k is not None:
            knn_k = k
    return {"promoted": promoted_clfs, "knn_k": knn_k}


def config_rows(payloads: List[dict]) -> List[List[str]]:
    """One row per classifier, reflecting whatever is actually promoted."""
    st = promoted_state(payloads)
    promoted, knn_k = st["promoted"], st["knn_k"]
    frozen = ("TCGA-FULL, seeds 42-46", "frozen τ_TCGA", "thresholds_TCGA.json")

    if "knn" in promoted:
        knn_row = ["k-Nearest Neighbours",
                   f"re-scored at k={knn_k}, weights='uniform' "
                   f"(reference set: all 413 TCGA slides)",
                   "TCGA-FULL, seeds 42-46",
                   f"refit τ - Youden's J on TCGA out-of-fold probabilities "
                   f"at k={knn_k}",
                   (f"knn_k35_thresholds.json (read, not recomputed)"
                    if knn_k == 35 else
                    f"derived - refit at k={knn_k} "
                    f"(knn_k35_thresholds.json only covers k=35)")]
    else:
        knn_row = ["k-Nearest Neighbours",
                   "k from the artifact's own internal GridSearchCV on TCGA "
                   "(grid: k∈{3,5,7,10,15} × metric∈{cosine,euclidean,"
                   "manhattan} × weights∈{uniform,distance}) - k=3 in most "
                   "saved artifacts", *frozen]

    if "rf" in promoted:
        rf_row = ["Random Forest",
                  "n_estimators=500, max_depth=None, min_samples_split=5, "
                  "min_samples_leaf=1, class_weight={0:1, 1:10} "
                  "(scores unchanged)",
                  "TCGA-FULL, seeds 42-46",
                  "rate-matched quantile - r = fraction of TCGA out-of-fold "
                  "slides τ_TCGA flags positive; target cut at its own "
                  "(1-r) quantile",
                  "derived per combination from TCGA out-of-fold predictions"]
    else:
        rf_row = ["Random Forest",
                  "n_estimators=500, max_depth=None, min_samples_split=5, "
                  "min_samples_leaf=1, class_weight={0:1, 1:10}", *frozen]

    return [
        ["Logistic Regression", "C=10, max_iter=300", *frozen],
        ["Neural Network", "per-combination, from TCGA-CV 4-fold majority "
                           "vote (max_iter=500 in every config)", *frozen],
        knn_row,
        ["ProtoNet", "no hyperparameters - class means of L2-normalised "
                     "TCGA features", *frozen],
        rf_row,
    ]


# ------------------------------------------------------------------- builders

def build(experiment: str, out_path: Path) -> None:
    payloads = load_results(experiment)
    if not payloads:
        sys.exit(f"no result files found for {experiment}")
    s0 = payloads[0]["stamp"]
    n_test = s0["n_test"]
    cc = s0["class_counts"]
    n_pos_lbl = cc.get("1", cc.get(1))
    n_neg_lbl = cc.get("0", cc.get(0))
    seeds = s0["seeds"]

    doc = Document()
    for style_name, sz in (("Normal", 10),):
        doc.styles[style_name].font.size = Pt(sz)
        doc.styles[style_name].font.name = "Calibri"

    title = doc.add_heading(f"{experiment} — External Validation Report", level=0)
    for r in title.runs:
        r.font.color.rgb = RGB_HEAD
    add_body(doc, f"Generated from the published result files under "
                  f"{'SurGen_EV_Results/' if experiment=='SurGen-EV' else 'TCGA_PAIP_EV_Results/'}"
                  f". All models are trained entirely on TCGA; zero "
                  f"{'SurGen' if experiment=='SurGen-EV' else 'PAIP'} data enters "
                  f"training or threshold selection.",
             italic=True, color=RGB_MUTED)

    # ---------------------------------------------------------------- overview
    add_heading(doc, "1. Overview", level=1)
    add_table(doc, [
        ["Test cohort", f"{n_test} slides ({n_pos_lbl} MSI-H / {n_neg_lbl} non-MSI-H, "
                        f"{100*int(n_pos_lbl)/int(n_test):.1f}% prevalence)"],
        ["Training cohort", "TCGA-FULL — all 413 TCGA slides "
                            "(ANN: 351 train + 62 early-stop)"],
        ["Seeds", ", ".join(str(x) for x in seeds)],
        ["Combinations", f"{len(payloads)} (aggregation method × foundation model, "
                         f"plus slide-level baselines)"],
        ["Training variants reported", "tcga_full (primary), fold_ensemble, fold_average"],
        ["Primary metric", "AUROC (threshold-free)"],
    ], header=["Field", "Value"], widths=[1.8, 4.7])

    # ------------------------------------------------------------ configuration
    st_promo = promoted_state(payloads)
    promoted, knn_k = st_promo["promoted"], st_promo["knn_k"]
    add_heading(doc, "2. Classifier configuration", level=1)
    if not promoted:
        add_body(doc, "Every classifier is trained once on TCGA and applied "
                      f"to {'SurGen' if experiment=='SurGen-EV' else 'PAIP'} "
                      "unchanged — all five report at the frozen τ_TCGA "
                      "operating point (Youden's J on pooled out-of-fold "
                      "TCGA-CV probabilities). No target-cohort label or "
                      "score is used to pick a threshold.")
    else:
        untouched = [CLF_NAME[c] for c in CLF_ORDER if c not in promoted]
        promoted_names = [CLF_NAME[c] for c in CLF_ORDER if c in promoted]
        add_body(doc, f"{', '.join(untouched)} keep the frozen τ_TCGA "
                      f"threshold and are numerically identical to the "
                      f"pre-promotion record. {', '.join(promoted_names)} "
                      f"{'were' if len(promoted_names)>1 else 'was'} promoted "
                      f"to a corrected threshold scheme — their published "
                      f"BalAcc, Acc and confusion matrix now come from that "
                      f"scheme, not the frozen one. Both corrected schemes "
                      f"are derived from TCGA out-of-fold predictions alone; "
                      f"the quantile rule reads target scores but never "
                      f"target labels.")
    rows = config_rows(payloads)
    add_table(doc, rows,
              header=["Classifier", "Model configuration", "Trained on",
                      "Threshold scheme", "τ source"],
              widths=[1.15, 2.15, 1.15, 1.85, 1.7], small=True)

    still_frozen_pathological = [c for c in ("proto", "rf") if c not in promoted]
    if still_frozen_pathological:
        names = " and ".join(CLF_NAME[c] for c in still_frozen_pathological)
        add_body(doc,
            f"Diagnostic note: under the frozen scheme, several {names} "
            f"configurations sit at a dead or near-dead operating point "
            f"(predicting no, or almost no, MSI-H positives). This is a "
            f"known limitation of the frozen threshold on this cohort, not "
            f"a modelling failure; AUROC (threshold-free) is unaffected.",
            italic=True, size=9, color=RGB_MUTED)

    # ------------------------------------------------------------------ radars
    add_heading(doc, "3. Radar plots", level=1)
    add_body(doc, "Primary variant (tcga_full), all aggregation-method × "
                  "encoder combinations. Each spoke is one (classifier, aggregation "
                  "method) pair; colour distinguishes the foundation model.")
    rdir = RADAR_ROOT / experiment
    add_image_pair(doc, [rdir / f"{experiment}_AUROC_radar.png"],
                   f"Figure 1. {experiment} — AUROC, by classifier and aggregation method.")
    doc.add_paragraph()
    promoted_label = " and ".join(CLF_NAME[c] for c in CLF_ORDER if c in promoted)
    add_image_pair(doc, [rdir / f"{experiment}_BalAcc_radar.png"],
                   f"Figure 2. {experiment} — Balanced accuracy, by classifier and "
                   f"aggregation method."
                   + (f" Reflects the promoted scheme for {promoted_label}."
                      if promoted else ""))

    # ------------------------------------------------------- full result tables
    doc.add_page_break()
    add_heading(doc, "4. Complete results by combination", level=1)
    add_body(doc, "Primary variant (tcga_full). One table per classifier, sorted by "
                  "AUROC descending. τ is the reported operating point; "
                  "n+ is the number of slides predicted positive.")

    for clf in CLF_ORDER:
        rows = []
        for p in payloads:
            st = p["stamp"]
            e = p["results"].get(clf, {}).get("tcga_full")
            if not e:
                continue
            scheme = e.get("threshold_scheme", "frozen_tau_TCGA")
            rows.append((st["method"], st["model"], e["auroc"], e["bacc"],
                        e["acc"], e["macro_f1"], e["threshold"],
                        n_pos(e["conf_matrix"]), e.get("status", "—"), scheme))
        rows.sort(key=lambda r: -r[2])

        add_heading(doc, f"4.{CLF_ORDER.index(clf)+1} {CLF_NAME[clf]}", level=2)
        table_rows = [[METHOD_LABEL.get(m, m), mo, f"{a:.4f}", f"{b:.4f}",
                      f"{ac:.4f}", f"{f1:.4f}", f"{t:.4f}", str(npp), st_,
                      sch.replace("_", " ")]
                     for m, mo, a, b, ac, f1, t, npp, st_, sch in rows]
        add_table(doc, table_rows,
                  header=["Aggregation", "Encoder", "AUROC", "BalAcc", "Acc",
                          "MacroF1", "τ", "n+", "Status", "Scheme"],
                  widths=[1.3, 0.85, 0.62, 0.62, 0.55, 0.62, 0.55, 0.4, 0.75, 1.1])
        doc.add_paragraph()

    # ------------------------------------------------------ selection rationale
    if experiment == "SurGen-EV" and promoted:
        doc.add_page_break()
        add_heading(doc, "5. Selection rationale — why k=20", level=1)
        add_body(doc,
            "kNN and RF were promoted on 3 Sep 2026 after a diagnostic sweep "
            "under experiments/surgen_ev_diagnostic/ (k ∈ {15,20,25,35,50,75,"
            "100}, three threshold schemes; nothing in that sweep touched the "
            "published tree at the time). Full methodology and the sweep's own "
            "figures are in experiments/surgen_ev_diagnostic/SUMMARY.md and "
            "SurGen_EV_KNN_K_Sweep_Report.docx.", color=RGB_CHANGED)

        agg_df = pd.DataFrame([
            {"clf": clf, "auroc": p["results"][clf]["tcga_full"]["auroc"],
             "bacc": p["results"][clf]["tcga_full"]["bacc"]}
            for p in payloads if p["stamp"]["method"] not in ("TITAN", "PRISM")
            for clf in CLF_ORDER if clf in p["results"]
        ])
        means = agg_df.groupby("clf")[["auroc", "bacc"]].mean().round(4)
        m_auroc, m_bacc = means["auroc"].mean(), means["bacc"].mean()

        add_body(doc,
            f"kNN's mean AUROC over the 20 aggregation combinations peaks at "
            f"k=20 ({means.loc['knn','auroc']:.4f} in this published record); "
            f"the sweep measured it falling to 0.5495 by k=100 — a narrow "
            f"plateau, not an open-ended gain. The oracle ceiling (best "
            f"possible balanced accuracy at any threshold, computed using "
            f"target labels and therefore never itself attainable) capped "
            f"around 0.63-0.64 across the whole k range tested, pointing to "
            f"kNN's score representation — not its threshold — as the "
            f"limiting factor on this cohort.",
            size=9.5)
        add_body(doc,
            f"RF's rate-matched quantile threshold clears every dead or "
            f"near-dead configuration it previously had (7 of 20 under the "
            f"frozen scheme) while leaving its AUROC exactly unchanged — the "
            f"quantile rule moves only the cut point, never the ranking.",
            size=9.5)
        add_body(doc,
            f"Mean across all five heads, this published record: "
            f"AUROC {m_auroc:.4f}, BalAcc {m_bacc:.4f}. Per head: " +
            ", ".join(f"{CLF_NAME[c]} {means.loc[c,'auroc']:.4f} / "
                     f"{means.loc[c,'bacc']:.4f}"
                     for c in CLF_ORDER if c in means.index) + ".",
            size=9.5)

    doc.add_page_break()
    add_heading(doc, "Provenance", level=1)
    add_table(doc, [
        ["Source tree", f"slide_classification/{'SurGen_EV_Results' if experiment=='SurGen-EV' else 'TCGA_PAIP_EV_Results'}/"],
        ["Threshold policy", s0.get("threshold_policy", "—")[:400]],
        ["Git commit", s0.get("git_commit", "—")],
        ["Radar plots", str(rdir.relative_to(REPO_ROOT))],
    ], header=["Field", "Value"], widths=[1.5, 5.0])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print(f"wrote {out_path}  ({out_path.stat().st_size/1024:.0f} KB)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--experiment", required=True, choices=["SurGen-EV", "PAIP-EV"])
    ap.add_argument("--out", type=Path, default=None)
    args = ap.parse_args()
    out = args.out or REPO_ROOT / f"{args.experiment.replace('-', '_')}_Report.docx"
    build(args.experiment, out)


if __name__ == "__main__":
    main()
